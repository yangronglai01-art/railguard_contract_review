"""合同审核FastAPI接口集成测试。"""

from collections.abc import Iterator
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient
from langgraph.checkpoint.memory import InMemorySaver

from railguard.api.main import create_app
from railguard.config import Settings
from railguard.models.schemas import Evidence
from railguard.rag.client import RagUnavailableError
from railguard.rag.models import RagSearchRequest
from railguard.storage.contracts import ContractRepository
from railguard.workflow.graph import build_review_graph
from railguard.workflow.service import ReviewService


def create_demo_docx_bytes() -> bytes:
    """在内存中创建会触发六项风险的演示合同。"""
    buffer = BytesIO()
    document = Document()

    document.add_paragraph(
        "设备监测平台软件采购合同"
    )
    document.add_paragraph(
        "第一条 项目内容"
    )
    document.add_paragraph(
        "乙方为甲方建设设备监测平台。"
    )
    document.add_paragraph(
        "第二条 付款方式"
    )
    document.add_paragraph(
        "合同签订后五日内，甲方支付全部合同款。"
    )
    document.add_paragraph(
        "第三条 验收"
    )
    document.add_paragraph(
        "系统上线三日后，甲方未提出异议视为验收合格。"
    )

    document.save(buffer)
    return buffer.getvalue()


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    """创建使用临时业务数据库和checkpoint的测试客户端。"""
    repository = ContractRepository(
        tmp_path / "review-api-contracts.db"
    )
    settings = Settings(
        rag_mode="mock",
        rag_mock_corpus_path=Path(
            "data/demo/rag-corpus.json"
        ),
        checkpoint_path=(
            tmp_path
            / "review-api-checkpoints.sqlite3"
        ),
    )
    application = create_app(
        repository,
        settings=settings,
    )

    with TestClient(application) as test_client:
        yield test_client


def upload_demo_contract(
    client: TestClient,
) -> dict[str, object]:
    """上传演示合同并返回接口JSON。"""
    response = client.post(
        "/contracts",
        files={
            "file": (
                "software-purchase-demo.docx",
                create_demo_docx_bytes(),
                (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
            )
        },
    )

    assert response.status_code == 201
    return response.json()


def test_review_api_starts_queries_and_resumes(
    client: TestClient,
) -> None:
    """验证审核接口从启动到人工完成的完整链路。"""
    contract = upload_demo_contract(client)
    contract_id = contract["contract_id"]

    start_response = client.post(
        f"/contracts/{contract_id}/reviews"
    )

    assert start_response.status_code == 201

    waiting = start_response.json()
    review_id = waiting["review_id"]

    assert waiting["contract_id"] == contract_id
    assert waiting["status"] == "awaiting_human"
    assert len(waiting["findings"]) == 6
    assert waiting["final_findings"] is None
    assert waiting["evidence"]

    get_response = client.get(
        f"/reviews/{review_id}"
    )

    assert get_response.status_code == 200
    assert get_response.json() == waiting

    # 未知风险ID应返回422，而且不能消耗人工审批中断。
    invalid_response = client.post(
        f"/reviews/{review_id}/decision",
        json={
            "action": "approve",
            "reviewer": "法务审核员",
            "finding_decisions": {
                "unknown-finding": "dismiss"
            },
            "comment": "测试无效风险ID。",
        },
    )

    assert invalid_response.status_code == 422

    still_waiting = client.get(
        f"/reviews/{review_id}"
    ).json()
    assert still_waiting["status"] == "awaiting_human"
    assert still_waiting["human_decision"] is None

    dismissed_finding_id = (
        waiting["findings"][0]["finding_id"]
    )
    decision_payload = {
        "action": "approve",
        "reviewer": "法务审核员",
        "finding_decisions": {
            dismissed_finding_id: "dismiss"
        },
        "comment": "第一项风险经核对后驳回。",
    }

    decision_response = client.post(
        f"/reviews/{review_id}/decision",
        json=decision_payload,
    )

    assert decision_response.status_code == 200

    completed = decision_response.json()

    assert completed["status"] == "approved"
    assert len(completed["findings"]) == 6
    assert len(completed["final_findings"]) == 5
    assert completed["human_decision"] == decision_payload
    assert "最终保留5项" in completed["final_summary"]

    # 已完成的审核不能重复提交人工决定。
    duplicate_response = client.post(
        f"/reviews/{review_id}/decision",
        json=decision_payload,
    )

    assert duplicate_response.status_code == 409


def test_review_api_returns_not_found(
    client: TestClient,
) -> None:
    """验证不存在的合同和审核任务返回404。"""
    missing_contract = client.post(
        "/contracts/not-found/reviews"
    )
    missing_review = client.get(
        "/reviews/not-found"
    )
    missing_decision = client.post(
        "/reviews/not-found/decision",
        json={
            "action": "approve",
            "reviewer": "法务审核员",
            "finding_decisions": {},
            "comment": "",
        },
    )

    assert missing_contract.status_code == 404
    assert missing_review.status_code == 404
    assert missing_decision.status_code == 404


class FailingRetriever:
    """始终模拟RAG连接失败的API测试检索器。"""

    async def retrieve(
        self,
        request: RagSearchRequest,
    ) -> list[Evidence]:
        """抛出RAG服务不可用异常。"""
        del request

        raise RagUnavailableError(
            "mock RAG service is unavailable"
        )


def test_review_api_returns_structured_execution_error(
    tmp_path: Path,
) -> None:
    """验证启动失败时返回可查询的review_id和稳定错误码。"""
    repository = ContractRepository(
        tmp_path / "failed-api-contracts.db"
    )
    graph = build_review_graph(
        retriever=FailingRetriever(),
        checkpointer=InMemorySaver(),
    )
    service = ReviewService(graph)
    settings = Settings(
        rag_mode="mock",
        checkpoint_path=(
            tmp_path
            / "unused-checkpoints.sqlite3"
        ),
    )
    application = create_app(
        repository,
        review_service=service,
        settings=settings,
    )

    with TestClient(application) as test_client:
        contract = upload_demo_contract(test_client)

        response = test_client.post(
            f"/contracts/{contract['contract_id']}/reviews"
        )

        assert response.status_code == 503

        detail = response.json()["detail"]

        assert detail["code"] == "rag_unavailable"
        assert detail["retryable"] is True
        assert detail["review_id"]

        failed_response = test_client.get(
            f"/reviews/{detail['review_id']}"
        )

        assert failed_response.status_code == 200
        assert failed_response.json()["status"] == "failed"