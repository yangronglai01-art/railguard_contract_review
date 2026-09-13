"""合同上传与查询API测试。"""

from collections.abc import Iterator
from io import BytesIO
from pathlib import Path

import pytest
from docx import Document
from fastapi.testclient import TestClient

from railguard.api.main import create_app
from railguard.parsers.documents import DEFAULT_MAX_FILE_SIZE
from railguard.storage.contracts import ContractRepository


def create_docx_bytes() -> bytes:
    """在内存中创建用于上传测试的DOCX合同。"""
    buffer = BytesIO()
    document = Document()
    document.add_paragraph("设备监测平台软件采购合同")
    document.add_paragraph("甲方验收合格后支付尾款。")
    document.save(buffer)

    return buffer.getvalue()


@pytest.fixture
def client(tmp_path: Path) -> Iterator[TestClient]:
    """创建使用临时SQLite数据库的测试客户端。

    测试结束后，pytest会删除临时数据库，
    不会修改data/runtime下的演示数据库。
    """
    repository = ContractRepository(
        tmp_path / "api-contracts.db"
    )
    application = create_app(repository)

    # 上下文管理器会执行FastAPI生命周期函数。
    with TestClient(application) as test_client:
        yield test_client


def test_upload_and_get_docx_contract(
    client: TestClient,
) -> None:
    """验证DOCX可以上传、保存并通过ID重新读取。"""
    upload_response = client.post(
        "/contracts",
        files={
            "file": (
                "monitoring-platform.docx",
                create_docx_bytes(),
                (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
            )
        },
    )

    assert upload_response.status_code == 201

    uploaded = upload_response.json()
    contract_id = uploaded["contract_id"]

    assert uploaded["filename"] == "monitoring-platform.docx"
    assert "设备监测平台软件采购合同" in uploaded["full_text"]
    # 测试合同没有“第X条”标题，因此回退为一个全文条款。
    assert len(uploaded["clauses"]) == 1
    assert uploaded["clauses"][0]["title"] == "合同全文"
    assert uploaded["clauses"][0]["text"] == uploaded["full_text"]

    get_response = client.get(
        f"/contracts/{contract_id}"
    )

    assert get_response.status_code == 200
    assert get_response.json() == uploaded


def test_reject_unsupported_upload(
    client: TestClient,
) -> None:
    """验证不支持的文件类型返回HTTP 415。"""
    response = client.post(
        "/contracts",
        files={
            "file": (
                "contract.txt",
                b"plain text",
                "text/plain",
            )
        },
    )

    assert response.status_code == 415


def test_reject_oversized_upload(
    client: TestClient,
) -> None:
    """验证超过10MB的上传返回HTTP 413。"""
    response = client.post(
        "/contracts",
        files={
            "file": (
                "contract.docx",
                b"x" * (DEFAULT_MAX_FILE_SIZE + 1),
                (
                    "application/vnd.openxmlformats-officedocument."
                    "wordprocessingml.document"
                ),
            )
        },
    )

    assert response.status_code == 413


def test_missing_contract_returns_404(
    client: TestClient,
) -> None:
    """验证不存在的合同ID返回HTTP 404。"""
    response = client.get("/contracts/not-found")

    assert response.status_code == 404
    assert response.json()["detail"] == "Contract not found."