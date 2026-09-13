"""DOCX和PDF合同解析测试。"""

from io import BytesIO

import pymupdf
import pytest
from docx import Document

from railguard.parsers.documents import (
    EmptyDocumentError,
    InvalidDocumentError,
    UnsupportedFileTypeError,
    normalize_text,
    parse_document,
)


def create_docx_bytes() -> bytes:
    """在内存中创建包含段落和表格的真实DOCX文件。"""
    buffer = BytesIO()
    document = Document()

    document.add_paragraph("软件采购合同")

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "付款条件"
    table.cell(0, 1).text = "验收合格后付款"

    document.save(buffer)
    return buffer.getvalue()


def create_pdf_bytes() -> bytes:
    """在内存中创建带文本层的真实PDF文件。"""
    document = pymupdf.open()
    page = document.new_page()

    page.insert_text(
        (72, 72),
        "Payment after acceptance.",
    )

    content = document.tobytes()
    document.close()
    return content


def test_normalize_text_compacts_blank_lines() -> None:
    """验证换行和连续空行能够被统一处理。"""
    original = "第一条  \r\n\r\n\r\n第二条\t\r\n"

    assert normalize_text(original) == "第一条\n\n第二条"


def test_parse_docx_includes_paragraphs_and_tables() -> None:
    """验证DOCX段落和表格内容均能被提取。"""
    text = parse_document(
        "software-contract.docx",
        create_docx_bytes(),
    )

    assert "软件采购合同" in text
    assert "付款条件\t验收合格后付款" in text


def test_parse_text_pdf() -> None:
    """验证带文本层的PDF可以正常提取文字。"""
    text = parse_document(
        "software-contract.pdf",
        create_pdf_bytes(),
    )

    assert "Payment after acceptance." in text


def test_reject_unsupported_extension() -> None:
    """验证不支持的文件扩展名会被拒绝。"""
    with pytest.raises(UnsupportedFileTypeError):
        parse_document("contract.txt", b"contract")


def test_reject_fake_pdf() -> None:
    """验证只修改扩展名的伪PDF会被拒绝。"""
    with pytest.raises(InvalidDocumentError):
        parse_document("contract.pdf", b"not a pdf")


def test_reject_empty_docx() -> None:
    """验证不包含有效文本的DOCX不能进入审核流程。"""
    buffer = BytesIO()
    Document().save(buffer)

    with pytest.raises(EmptyDocumentError):
        parse_document("empty.docx", buffer.getvalue())